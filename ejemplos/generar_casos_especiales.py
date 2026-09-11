"""
generar_casos_especiales.py
===========================
Genera casos contables especiales para practicar detección avanzada de errores.

Crea:
  - caso_duplicados.xlsx      (Mismo movimiento duplicado)
  - caso_valores_faltantes.xlsx (Cuentas/centros vacíos)
  - caso_tolerancia.xlsx      (Diferencia por redondeo)
  - caso_conciliacion.xlsx    (Movimiento faltante en una empresa)

Ejecutar con:
    python generar_casos_especiales.py
"""

import os
import pandas as pd
import numpy as np
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

PERIODO = "2024-08"
DIRECTORIO = os.path.dirname(os.path.abspath(__file__))

# Cuentas y centros de costo estándar
CUENTAS = {
    "5100001": "Remuneraciones Sueldo Base",
    "5100002": "Remuneraciones Horas Extras",
    "2105001": "AFP por Pagar",
    "1101001": "Banco Estado CTA CTE",
}

CENTROS_COSTO = ["CC001-Admin", "CC002-Operaciones", "CC003-Ventas"]


def generar_encabezado_sofland(ws, nombre: str):
    """Agrega las 4 filas de encabezado Sofland estándar."""
    ws["A1"] = "SISTEMA SOFLAND S.A."
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = f"Fecha Impresión: {date.today().strftime('%d/%m/%Y')}"
    ws["A3"] = f"Usuario: JOCELYN.CONTABILIDAD"
    ws["A4"] = "-" * 50
    return ws


def guardar_con_encabezado_sofland(df: pd.DataFrame, ruta: str, nombre: str):
    """Guarda DataFrame con encabezado Sofland de 4 filas."""
    with pd.ExcelWriter(ruta, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Comprobante", startrow=4, index=False)
        ws = writer.sheets["Comprobante"]
        generar_encabezado_sofland(ws, nombre)
    print(f"  ✅ {os.path.basename(ruta)}")


# ──────────────────────────────────────────────────────────────

def caso_duplicados():
    """Caso: Mismo movimiento duplicado (error común en importación)."""
    registros = [
        {"Cuenta": "5100001", "Glosa": "Sueldo base agosto", "Debe": 5000000.0, "Haber": 0.0, "CentroCosto": "CC001-Admin"},
        {"Cuenta": "5100001", "Glosa": "Sueldo base agosto", "Debe": 5000000.0, "Haber": 0.0, "CentroCosto": "CC001-Admin"},  # Duplicado
        {"Cuenta": "5100002", "Glosa": "Horas extras agosto", "Debe": 800000.0, "Haber": 0.0, "CentroCosto": "CC002-Operaciones"},
        {"Cuenta": "2105001", "Glosa": "AFP retenida agosto", "Debe": 0.0, "Haber": 580000.0, "CentroCosto": None},
        {"Cuenta": "1101001", "Glosa": "Pago nómina masiva", "Debe": 0.0, "Haber": 5220000.0, "CentroCosto": None},
    ]
    df = pd.DataFrame(registros)
    ruta = os.path.join(DIRECTORIO, "caso_duplicados.xlsx")
    guardar_con_encabezado_sofland(df, ruta, "CasoDuplicados")
    return df


def caso_valores_faltantes():
    """Caso: Cuenta vacía, centro de costo faltante, glosa vacía."""
    registros = [
        {"Cuenta": "5100001", "Glosa": "Sueldo base agosto", "Debe": 5000000.0, "Haber": 0.0, "CentroCosto": "CC001-Admin"},
        {"Cuenta": None, "Glosa": "Movimiento huérfano", "Debe": 1000000.0, "Haber": 0.0, "CentroCosto": "CC003-Ventas"},  # Cuenta vacía
        {"Cuenta": "5100002", "Glosa": None, "Debe": 800000.0, "Haber": 0.0, "CentroCosto": "CC002-Operaciones"},  # Glosa vacía
        {"Cuenta": "2105001", "Glosa": "AFP retenida agosto", "Debe": 0.0, "Haber": 580000.0, "CentroCosto": None},  # Centro de costo vacío
        {"Cuenta": "1101001", "Glosa": "Pago nómina", "Debe": 0.0, "Haber": 6220000.0, "CentroCosto": None},
    ]
    df = pd.DataFrame(registros)
    ruta = os.path.join(DIRECTORIO, "caso_valores_faltantes.xlsx")
    guardar_con_encabezado_sofland(df, ruta, "CasoValoresFaltantes")
    return df


def caso_tolerancia():
    """Caso: Descuadre menor por diferencia de redondeo (±1000 pesos)."""
    registros = [
        {"Cuenta": "5100001", "Glosa": "Sueldo base agosto", "Debe": 5000000.0, "Haber": 0.0, "CentroCosto": "CC001-Admin"},
        {"Cuenta": "5100002", "Glosa": "Horas extras agosto", "Debe": 800000.0, "Haber": 0.0, "CentroCosto": "CC002-Operaciones"},
        {"Cuenta": "2105001", "Glosa": "AFP retenida agosto", "Debe": 0.0, "Haber": 580000.0, "CentroCosto": None},
        {"Cuenta": "1101001", "Glosa": "Pago nómina masiva", "Debe": 0.0, "Haber": 5219000.0, "CentroCosto": None},  # Diferencia: 1000
    ]
    df = pd.DataFrame(registros)
    ruta = os.path.join(DIRECTORIO, "caso_tolerancia.xlsx")
    guardar_con_encabezado_sofland(df, ruta, "CasoTolerancia")
    return df


def caso_conciliacion():
    """Caso: Empresa con un movimiento faltante que la descuadra."""
    registros = [
        {"Cuenta": "5100001", "Glosa": "Sueldo base agosto", "Debe": 5000000.0, "Haber": 0.0, "CentroCosto": "CC001-Admin"},
        {"Cuenta": "5100002", "Glosa": "Horas extras agosto", "Debe": 800000.0, "Haber": 0.0, "CentroCosto": "CC002-Operaciones"},
        {"Cuenta": "2105001", "Glosa": "AFP retenida agosto", "Debe": 0.0, "Haber": 580000.0, "CentroCosto": None},
        # Falta: {"Cuenta": "1101001", "Glosa": "Pago nómina", "Debe": 0.0, "Haber": 5220000.0}
        # Descuadre: Debe=5800000, Haber=580000
    ]
    df = pd.DataFrame(registros)
    ruta = os.path.join(DIRECTORIO, "caso_conciliacion.xlsx")
    guardar_con_encabezado_sofland(df, ruta, "CasoConciliacion")
    return df


if __name__ == "__main__":
    print("=" * 60)
    print("  GENERADOR DE CASOS ESPECIALES")
    print(f"  Período: {PERIODO}")
    print("=" * 60)
    print()

    caso_duplicados()
    caso_valores_faltantes()
    caso_tolerancia()
    caso_conciliacion()

    print()
    print("  📁 Casos especiales creados en 'ejemplos/'")
    print()
    print("  GUÍA PARA PRACTICAR:")
    print("  1. caso_duplicados.xlsx → Detecta filas idénticas")
    print("  2. caso_valores_faltantes.xlsx → Detecta cuentas/glosas/centros vacíos")
    print("  3. caso_tolerancia.xlsx → Diferencia por redondeo (±1000)")
    print("  4. caso_conciliacion.xlsx → Movimiento faltante")
    print()
